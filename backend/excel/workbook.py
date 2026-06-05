from pathlib import Path
from openpyxl import Workbook, load_workbook

DATA_DIR = Path("excel")
FILE_PATH = DATA_DIR / "bautro-clinic.xlsx"

PATIENT_HEADERS = [
    "patient_number",
    "date_created",
    "full_name",
    "address",
    "contact_number",
    "date_of_birth",
    "age",
    "sex",
    "occupation",
    "emergency_contact",
    "chief_complaint",
    "medical_history",
    "ocular_history",
    "allergies",
    "current_medications",
    "visual_acuity_od",
    "visual_acuity_os",
    "bcva_od",
    "bcva_os",
    "iop",
    "diagnosis",
    "clinical_notes",
    "other_tests"
]

PRESCRIPTION_HEADERS = [
    "prescription_number",
    "patient_number",
    "date",

    "od_sphere",
    "od_cylinder",
    "od_axis",
    "od_add",

    "os_sphere",
    "os_cylinder",
    "os_axis",
    "os_add",

    "distance_pd",
    "near_pd",

    "right_seg_height",
    "left_seg_height",

    "right_pupil_height",
    "left_pupil_height",

    "lens_type",
    "lens_brand",
    "lens_coating",
    "tint",

    "frame_model",
    "frame_brand",

    "special_instructions"
]

INVENTORY_HEADERS = [
    "item_code",
    "category",
    "brand",
    "model",
    "color",
    "supplier",
    "purchase_cost",
    "selling_price",
    "quantity",
    "reorder_level"
]

BILLING_HEADERS = [
    "invoice_number",
    "patient_number",
    "product_code",
    "quantity",
    "unit_price",
    "total",
    "payment_method",
    "date"
]

INVENTORY_HISTORY_HEADERS = [
    "date",
    "item_code",
    "action",
    "quantity",
    "remarks"
]


def create_sheet_with_headers(wb, name, headers):
    ws = wb.create_sheet(name)
    ws.append(headers)


def init_workbook():
    DATA_DIR.mkdir(exist_ok=True)

    if FILE_PATH.exists():
        return

    wb = Workbook()

    default_sheet = wb.active
    wb.remove(default_sheet)

    create_sheet_with_headers(
        wb,
        "Patients",
        PATIENT_HEADERS
    )

    create_sheet_with_headers(
        wb,
        "Prescriptions",
        PRESCRIPTION_HEADERS
    )

    create_sheet_with_headers(
        wb,
        "Inventory",
        INVENTORY_HEADERS
    )

    create_sheet_with_headers(
        wb,
        "Billing",
        BILLING_HEADERS
    )

    create_sheet_with_headers(
        wb,
        "InventoryHistory",
        INVENTORY_HISTORY_HEADERS
    )

    wb.save(FILE_PATH)


def get_workbook():
    init_workbook()
    return load_workbook(FILE_PATH)


def save_workbook(workbook):
    workbook.save(FILE_PATH)